from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.oxml.ns import qn

prs = Presentation('/home/user/Business-Modeling/Tesla_Filled.pptx')
s6 = prs.slides[5]

FONT  = '맑은 고딕'
RED   = RGBColor(0xE3, 0x19, 0x37)
DARK  = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY = RGBColor(0xF0, 0xF0, 0xF0)
DGRAY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x1E, 0x8B, 0x4C)
BLUE  = RGBColor(0x16, 0x4F, 0x96)
SLATE = RGBColor(0x2C, 0x3E, 0x50)
GOLD  = RGBColor(0xFF, 0xCC, 0x00)

# ── 기존에 추가된 도형 모두 제거 (원래 3개만 남기기) ──
keep_names = {'직사각형 1', '직선 연결선 3', 'Title 1'}
to_remove = [sh for sh in s6.shapes if sh.name not in keep_names]
for sh in reversed(to_remove):
    sp = sh._element
    sp.getparent().remove(sp)

# ── 헬퍼 ───────────────────────────────────────────────
def add_rect(slide, l, t, w, h, bg):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = bg
    sh.line.fill.background()
    return sh

def set_para(tf, items, clear=True):
    """items: list of (text, size, bold, color, align)"""
    if clear:
        for p in tf._txBody.findall(qn('a:p')):
            tf._txBody.remove(p)
    tf.word_wrap = True
    first = True
    for text, size, bold, color, align in items:
        if first:
            p = tf.paragraphs[0] if tf.paragraphs else tf.add_paragraph()
            first = False
        else:
            p = tf.add_paragraph()
        p.text = text
        p.alignment = align
        p.space_before = Pt(2)
        p.space_after  = Pt(1)
        r = p.runs[0] if p.runs else p.add_run()
        r.font.name  = FONT
        r.font.size  = Pt(size)
        r.font.bold  = bold
        r.font.color.rgb = color

# ── 레이아웃 기준 ───────────────────────────────────────
LX = 450000          # left margin
TY = 2080000         # top of content area
W  = 8200000         # total content width
LW = 3900000         # left column width
RW = W - LW - 80000 # right column width
RX = LX + LW + 80000

# ═══════════════════════════════════════════════════════
# 왼쪽 컬럼: 재무 현황 + 시나리오 A/B
# ═══════════════════════════════════════════════════════

# ── 섹션 헤더: 재무 현황 ────────────────────────────────
h = add_rect(s6, LX, TY, LW, 340000, RED)
set_para(h.text_frame, [("📊  재무 현황 (2023 Q3)", 13, True, WHITE, PP_ALIGN.LEFT)])
h.text_frame.paragraphs[0].space_before = Pt(6)

# ── 재무 항목 3개 나열 ──────────────────────────────────
kpi_box = add_rect(s6, LX, TY + 360000, LW, 700000, LGRAY)
set_para(kpi_box.text_frame, [
    ("영업이익률   17.2%  →  7.6%  ↓", 12, True,  RED,   PP_ALIGN.LEFT),
    ("순이익 증감  YoY −44%",           12, True,  RED,   PP_ALIGN.LEFT),
    ("美 시장점유율  62%  →  50%  ↓",   12, True,  RED,   PP_ALIGN.LEFT),
])
kpi_box.text_frame.paragraphs[0].space_before = Pt(10)

# ── 섹션 헤더: 시나리오 ────────────────────────────────
h2 = add_rect(s6, LX, TY + 1120000, LW, 340000, SLATE)
set_para(h2.text_frame, [("⚖️  두 가지 시나리오", 13, True, WHITE, PP_ALIGN.LEFT)])
h2.text_frame.paragraphs[0].space_before = Pt(6)

# ── Scenario A ─────────────────────────────────────────
sa = add_rect(s6, LX, TY + 1520000, LW, 860000, GREEN)
set_para(sa.text_frame, [
    ("✓  Scenario A  |  성공 경로",                         12, True,  WHITE,                PP_ALIGN.LEFT),
    ("볼륨↑  →  배터리 원가↓",                              11, False, RGBColor(0xCC,0xFF,0xCC), PP_ALIGN.LEFT),
    ("FSD 구독 시장 확대  →  고마진 반복 수익 확보",          11, False, RGBColor(0xCC,0xFF,0xCC), PP_ALIGN.LEFT),
    ("소프트웨어·에너지 수익으로 차량 마진 하락 상쇄",        11, False, RGBColor(0xCC,0xFF,0xCC), PP_ALIGN.LEFT),
])
sa.text_frame.paragraphs[0].space_before = Pt(8)

# ── Scenario B ─────────────────────────────────────────
sb = add_rect(s6, LX, TY + 2440000, LW, 860000, RED)
set_para(sb.text_frame, [
    ("✗  Scenario B  |  위험 경로",                         12, True,  WHITE,                PP_ALIGN.LEFT),
    ("브랜드 희석  →  기존 고소득 고객 이탈",                11, False, RGBColor(0xFF,0xDD,0xDD), PP_ALIGN.LEFT),
    ("ASP(평균판매가) 하락  →  마진 구조 붕괴",              11, False, RGBColor(0xFF,0xDD,0xDD), PP_ALIGN.LEFT),
    ("BYD식 마진 4~5% 구조 추종 시 수익성 모델 붕괴",        11, False, RGBColor(0xFF,0xDD,0xDD), PP_ALIGN.LEFT),
])
sb.text_frame.paragraphs[0].space_before = Pt(8)

# ═══════════════════════════════════════════════════════
# 오른쪽 컬럼: 3가지 해자 (선두 유지 전략)
# ═══════════════════════════════════════════════════════

# ── 섹션 헤더 ──────────────────────────────────────────
h3 = add_rect(s6, RX, TY, RW, 340000, DARK)
set_para(h3.text_frame, [("🏰  선두 유지 전략 : 3가지 해자(Moat)", 13, True, WHITE, PP_ALIGN.LEFT)])
h3.text_frame.paragraphs[0].space_before = Pt(6)

# ── 해자 ① 기술 ───────────────────────────────────────
m1 = add_rect(s6, RX, TY + 360000, RW, 960000, BLUE)
set_para(m1.text_frame, [
    ("🔬  해자 ①  기술 우위",                              12, True,  WHITE,                  PP_ALIGN.LEFT),
    ("FSD 완성도 조기 확보",                               11, False, RGBColor(0xCC,0xDD,0xFF), PP_ALIGN.LEFT),
    ("4680 배터리 양산 가속  →  원가 경쟁력 확보",         11, False, RGBColor(0xCC,0xDD,0xFF), PP_ALIGN.LEFT),
    ("수백만 대 주행 데이터 독점  →  경쟁사 복제 불가",    11, False, RGBColor(0xCC,0xDD,0xFF), PP_ALIGN.LEFT),
])
m1.text_frame.paragraphs[0].space_before = Pt(8)

# ── 해자 ② 생태계 ─────────────────────────────────────
m2 = add_rect(s6, RX, TY + 1380000, RW, 960000, SLATE)
set_para(m2.text_frame, [
    ("⚡  해자 ②  생태계 우위",                            12, True,  WHITE,                  PP_ALIGN.LEFT),
    ("Supercharger  →  수익 인프라로 전환",                11, False, RGBColor(0xCC,0xEE,0xFF), PP_ALIGN.LEFT),
    ("에너지·차량·보험 번들  →  Lock-in 구조 강화",       11, False, RGBColor(0xCC,0xEE,0xFF), PP_ALIGN.LEFT),
    ("V2G 상용화  →  이탈 비용 급상승",                   11, False, RGBColor(0xCC,0xEE,0xFF), PP_ALIGN.LEFT),
])
m2.text_frame.paragraphs[0].space_before = Pt(8)

# ── 해자 ③ 브랜드 ─────────────────────────────────────
m3 = add_rect(s6, RX, TY + 2400000, RW, 960000, DGRAY)
set_para(m3.text_frame, [
    ("🌍  해자 ③  브랜드·미션 우위",                       12, True,  WHITE,                  PP_ALIGN.LEFT),
    ("대중 시장엔 '경제적 합리성' 메시지로 전환",          11, False, RGBColor(0xEE,0xEE,0xEE), PP_ALIGN.LEFT),
    ("형평성 내러티브 선점  (CSR이 아닌 핵심 전략)",       11, False, RGBColor(0xEE,0xEE,0xEE), PP_ALIGN.LEFT),
    ("CEO 리스크 분산  →  브랜드 독립성 강화",             11, False, RGBColor(0xEE,0xEE,0xEE), PP_ALIGN.LEFT),
])
m3.text_frame.paragraphs[0].space_before = Pt(8)

# ═══════════════════════════════════════════════════════
# 하단 결론 바 (전체 폭)
# ═══════════════════════════════════════════════════════
conc = add_rect(s6, LX, TY + 3420000, W, 340000, DARK)
set_para(conc.text_frame, [(
    '결론  ·  "차를 팔고 끝나는 모델"  →  에너지 · 소프트웨어 구독 플랫폼 회사로의 전환이 Tesla의 장기 생존 전략',
    12, True, GOLD, PP_ALIGN.CENTER
)])
conc.text_frame.paragraphs[0].space_before = Pt(8)

prs.save('/home/user/Business-Modeling/Tesla_Filled.pptx')
print("Done.")
