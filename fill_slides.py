from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

prs = Presentation('/root/.claude/uploads/e6e685f1-ecb8-5b5f-b6be-d01e6dec7bb8/d21dd043-62.Teslas_EvolutionMass_Market_Electric_VehiclesTemplate.pptx')

# ── Colors ──────────────────────────────────────────────
RED    = RGBColor(0xE3, 0x19, 0x37)  # Tesla Red
DARK   = RGBColor(0x1A, 0x1A, 0x1A)
DGRAY  = RGBColor(0x3D, 0x3D, 0x3D)
MGRAY  = RGBColor(0x70, 0x70, 0x70)
LGRAY  = RGBColor(0xF4, 0xF4, 0xF4)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BLUE   = RGBColor(0x16, 0x4F, 0x96)
GREEN  = RGBColor(0x27, 0x9E, 0x60)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
SLATE  = RGBColor(0x2C, 0x3E, 0x50)
FONT   = '맑은 고딕'

def rect(slide, l, t, w, h, bg, border=False, border_color=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    if border and border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    return shape

def styled_tf(shape, items, font=FONT):
    """items: list of (text, size, bold, color, align, space_before, space_after)"""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None

    first = True
    for (text, size, bold, color, align, sp_b, sp_a) in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = text
        p.alignment = align
        if sp_b is not None:
            p.space_before = Pt(sp_b)
        if sp_a is not None:
            p.space_after = Pt(sp_a)
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

def add_label_content(slide, shape, label, label_color, content_lines, content_color=None, label_size=11, content_size=9):
    """Replace shape text with label + content"""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None

    # Clear existing paragraphs
    from pptx.oxml.ns import qn
    txBody = tf._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)

    items = [(label, label_size, True, label_color, PP_ALIGN.LEFT, 0, 3)]
    for line in content_lines:
        items.append((line, content_size, False, content_color or WHITE, PP_ALIGN.LEFT, 1, 1))

    first = True
    for (text, size, bold, color, align, sp_b, sp_a) in items:
        from pptx.oxml import parse_xml
        from lxml import etree
        p_xml = '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
        p_elem = parse_xml(p_xml)
        txBody.append(p_elem)
        from pptx.text.text import _Paragraph
        # Use tf.add_paragraph approach

    # Simpler approach: use tf directly
    tf2 = shape.text_frame
    # Re-get after clear
    first2 = True
    for (text, size, bold, color, align, sp_b, sp_a) in items:
        if first2:
            p = tf2.paragraphs[0] if tf2.paragraphs else tf2.add_paragraph()
            first2 = False
        else:
            p = tf2.add_paragraph()
        p.text = text
        p.alignment = align
        if sp_b: p.space_before = Pt(sp_b)
        if sp_a: p.space_after = Pt(sp_a)
        if p.runs:
            r = p.runs[0]
        else:
            r = p.add_run()
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color


# ══════════════════════════════════════════════════════════
# SLIDE 4  ─  Q1: 마케팅 & 제품 전략
# ══════════════════════════════════════════════════════════
s4 = prs.slides[3]
Y0 = 2100000   # start below the line
M  = 914400 // 10  # 0.1" margin
W  = 8200000   # total usable width
LX = 450000    # left start x
COL = (W - M) // 2
H_LABEL = 370000
H_TIER  = 560000
H_CARD  = 480000

# ── LEFT: 제품 전략 header ──────────────────────────────
hdr_l = rect(s4, LX, Y0, COL - M//2, H_LABEL, RED)
styled_tf(hdr_l, [("▸ 제품 전략 : 계단식 확장", 13, True, WHITE, PP_ALIGN.LEFT, 4, 0)])

tier_colors = [RED, SLATE, BLUE]
tier_texts = [
    ("① Premium", "Model S / X · Cybertruck",    "고소득 테크 얼리어답터 — 브랜드 핵심 수호"),
    ("② Mid-Market", "Model 3 / Y",              "중산층 실용 프리미엄 — TCO 절감 메시지"),
    ("③ Mass Market", "보급형 $25K↓ (신규)",     "Apple SE 전략 — 작지만 완벽한 Tesla"),
]
for i, (tier, model, desc) in enumerate(tier_texts):
    ty = Y0 + H_LABEL + M//4 + i * (H_TIER + M//4)
    b = rect(s4, LX, ty, COL - M//2, H_TIER, tier_colors[i])
    styled_tf(b, [
        (tier,  12, True,  WHITE, PP_ALIGN.LEFT, 6, 0),
        (model, 10, False, RGBColor(0xFF,0xCC,0xCC) if i==0 else RGBColor(0xAA,0xCC,0xFF), PP_ALIGN.LEFT, 2, 0),
        (desc,   9, False, RGBColor(0xDD,0xDD,0xDD), PP_ALIGN.LEFT, 3, 0),
    ])

# ── RIGHT: 마케팅 전략 header ───────────────────────────
RX = LX + COL + M//2
hdr_r = rect(s4, RX, Y0, COL - M//2, H_LABEL, DARK)
styled_tf(hdr_r, [("▸ 마케팅 전략 : 4가지 실행 방향", 13, True, WHITE, PP_ALIGN.LEFT, 4, 0)])

cards = [
    (RED,   "① 세그먼트별 메시지 분리",
             "고소득 → 기술·혁신·지위재  /  중산층 → TCO 절감\n저소득 → IRA 보조금 $7,500 + 연료비 절약 강조"),
    (DGRAY, "② 유통 채널 다각화",
             "직영 온라인 + 쇼룸 유지\n농촌·교외 서비스센터 확대 — 충전 인프라 先 투자"),
    (DGRAY, "③ 금융 접근성 강화",
             "리스(Lease) 프로그램 / 인증 중고(CPO) 시장\nIRA 세액공제 POS 즉시 적용 시스템 구축"),
    (SLATE, "④ 커뮤니티 파트너십",
             "비영리단체·CDFI와 저소득 커뮤니티 접근\n지역 체험·시승 프로그램 → 신뢰 구축"),
]
for i, (bg, title, body) in enumerate(cards):
    cy = Y0 + H_LABEL + M//4 + i * (H_CARD + M//4)
    c = rect(s4, RX, cy, COL - M//2, H_CARD, bg)
    styled_tf(c, [
        (title, 11, True,  WHITE, PP_ALIGN.LEFT, 6, 2),
        (body,   9, False, RGBColor(0xDD,0xDD,0xDD), PP_ALIGN.LEFT, 2, 0),
    ])

# ── BOTTOM: 핵심 원칙 ───────────────────────────────────
by = Y0 + H_LABEL + M//4 + 3 * (H_TIER + M//4)
bh = 320000
b_box = rect(s4, LX, by + 80000, W - M, bh, RGBColor(0xF9,0xF2,0xEC), border=True, border_color=RED)
styled_tf(b_box, [(
    '핵심 원칙  ·  "누구나 갖고 싶은 차 중에 살 수 있는 것"  —  브랜드 희석 없는 계단식 시장 확장',
    11, True, RED, PP_ALIGN.CENTER, 8, 0
)])


# ══════════════════════════════════════════════════════════
# SLIDE 5  ─  Q2: 5가지 추가 역량
# ══════════════════════════════════════════════════════════
s5 = prs.slides[4]
Y0_5 = 1900000
LX5  = 430000
W5   = 8300000
GAP  = 80000

caps = [
    (RED,   "① 저비용\n제조 역량",
             "4680 배터리 양산 가속\n플랫폼 부품 수 최소화\n기가팩토리 수율 극대화\n→ 목표: $25K EV 수익성"),
    (SLATE, "② 충전 인프라\n확장",
             "농촌·교외 Supercharger\nNACS 개방으로 수익화\nV2G 상용화 추진\n무선충전 R&D 투자"),
    (BLUE,  "③ 금융·접근성\n서비스",
             "리스 프로그램 확대\n인증 중고(CPO) 강화\n저신용 파이낸싱\nIRA $7,500 즉시 적용"),
    (DGRAY, "④ 세입자 충전\n솔루션",
             "공동주택 충전 설치 지원\n부동산사 파트너십\n미국 인구 36% 세입자\n→ 미개척 최대 시장"),
    (GREEN, "⑤ 커뮤니티\n신뢰 구축",
             "교육·시승 프로그램\n다양성 반영 마케팅\nBlueHub Capital 협력\n현지 서비스센터 확대"),
]

CW = (W5 - 4 * GAP) // 5
CH5 = 4100000

for i, (bg, title, body) in enumerate(caps):
    cx = LX5 + i * (CW + GAP)
    c = rect(s5, cx, Y0_5, CW, CH5, bg)
    styled_tf(c, [
        (title, 12, True,  WHITE, PP_ALIGN.CENTER, 12, 6),
        (body,   9, False, RGBColor(0xEE,0xEE,0xEE), PP_ALIGN.LEFT, 6, 2),
    ])

# sub-label below
sub_y = Y0_5 + CH5 + 100000
sub = rect(s5, LX5, sub_y, W5 - GAP, 260000, LGRAY)
styled_tf(sub, [(
    "우선순위  ①저비용제조 ≥ ②충전인프라  >  ③금융접근성 ≥ ④세입자  >  ⑤커뮤니티  —  생태계 전체를 함께 구축",
    10, True, DGRAY, PP_ALIGN.CENTER, 6, 0
)])


# ══════════════════════════════════════════════════════════
# SLIDE 6  ─  Q3: 장기 지속가능성 & 선두 유지 전략
# ══════════════════════════════════════════════════════════
s6 = prs.slides[5]
Y0_6 = 2100000
LX6  = 430000
W6   = 8300000

# ── Row 1: KPI 3개 ─────────────────────────────────────
kpis = [
    ("영업이익률", "17.2%  →  7.6%  ↓", RED),
    ("순이익 증감", "YoY  −44%",          RED),
    ("美 시장점유율", "62%  →  50%  ↓",  RED),
]
KW = (W6 - 2 * GAP) // 3
KH = 520000
for i, (label, val, vc) in enumerate(kpis):
    kx = LX6 + i * (KW + GAP)
    box = rect(s6, kx, Y0_6, KW, KH, DARK)
    styled_tf(box, [
        (label, 10, False, RGBColor(0xAA,0xAA,0xAA), PP_ALIGN.CENTER, 6, 2),
        (val,   16, True,  WHITE, PP_ALIGN.CENTER, 0, 0),
    ])

# ── Row 2: Scenario A vs B ──────────────────────────────
row2_y = Y0_6 + KH + GAP
HW = (W6 - GAP) // 2
SH = 900000

sa = rect(s6, LX6, row2_y, HW, SH, GREEN)
styled_tf(sa, [
    ("✓ Scenario A  |  성공 경로", 12, True, WHITE, PP_ALIGN.LEFT, 8, 4),
    ("볼륨↑ → 배터리 원가↓ → FSD 구독 시장 확대\n→ 소프트웨어 고마진으로 차량 마진 하락 상쇄\n→ 에너지·보험·충전 번들 수익으로 지속 성장", 9, False, RGBColor(0xDD,0xFF,0xDD), PP_ALIGN.LEFT, 4, 0),
])

sb = rect(s6, LX6 + HW + GAP, row2_y, HW, SH, RED)
styled_tf(sb, [
    ("✗ Scenario B  |  위험 경로", 12, True, WHITE, PP_ALIGN.LEFT, 8, 4),
    ("브랜드 희석 → 기존 고소득 고객 이탈\n→ ASP(평균판매가) 하락 → 마진 구조 붕괴\n→ BYD 마진 4~5% 구조 추종 시 수익성 모델 붕괴", 9, False, RGBColor(0xFF,0xDD,0xDD), PP_ALIGN.LEFT, 4, 0),
])

# ── Row 3: 3개 해자 ────────────────────────────────────
row3_y = row2_y + SH + GAP
moat_colors = [BLUE, SLATE, DGRAY]
moats = [
    ("🔬 해자 ① 기술 우위",
     "FSD 완성도 조기 확보\n4680 배터리 양산 가속\n수백만 대 주행 데이터 독점\n→ 경쟁사 단기 복제 불가"),
    ("⚡ 해자 ② 생태계 우위",
     "Supercharger → 수익 인프라 전환\n에너지·차량·보험 번들 락인(Lock-in)\nV2G 상용화 → 이탈 비용 급상승\n→ Tesla 생태계 안이 더 유리"),
    ("🌍 해자 ③ 브랜드·미션",
     "대중 시장엔 '경제적 합리성' 메시지\n형평성 내러티브 선점 (CSR이 아닌 전략)\nCEO 리스크 분산 — 브랜드 독립성 강화\n→ 미션 중심 장기 충성도"),
]
MW = (W6 - 2 * GAP) // 3
MH = 1400000
for i, (title, body) in enumerate(moats):
    mx = LX6 + i * (MW + GAP)
    m = rect(s6, mx, row3_y, MW, MH, moat_colors[i])
    styled_tf(m, [
        (title, 11, True,  WHITE, PP_ALIGN.LEFT, 8, 4),
        (body,   9, False, RGBColor(0xDD,0xDD,0xEE), PP_ALIGN.LEFT, 4, 0),
    ])

# ── Bottom conclusion ───────────────────────────────────
conc_y = row3_y + MH + GAP
conc = rect(s6, LX6, conc_y, W6 - GAP, 300000, RGBColor(0xF9,0xF2,0xEC), border=True, border_color=RED)
styled_tf(conc, [(
    '결론  ·  "차를 팔고 끝나는 모델"  →  에너지·소프트웨어 구독 플랫폼 회사로의 전환이 장기 생존 전략',
    11, True, RED, PP_ALIGN.CENTER, 8, 0
)])


# ══════════════════════════════════════════════════════════
# SLIDE 7  ─  Q4: Business Model Canvas
# ══════════════════════════════════════════════════════════
s7 = prs.slides[6]

# Shape map by label text
bmc_map = {}
for shape in s7.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        bmc_map[t] = shape

bmc_content = {
    "KP": {
        "label": "KP  |  핵심 파트너",
        "lines": [
            "• 리튬·원자재 공급업체",
            "• 커뮤니티 금융기관 (CDFI)",
            "• 비영리단체 (BlueHub 등)",
            "• 부동산·아파트 운영사",
            "• 지자체·연방정부 (IRA)",
            "• 무선충전 파트너 (WiTricity)",
        ],
        "label_color": WHITE,
        "bg": SLATE,
    },
    "KA": {
        "label": "KA  |  핵심 활동",
        "lines": [
            "• EV 설계·제조 (수직계열화)",
            "• 배터리 R&D (4680 셀)",
            "• FSD 소프트웨어 개발",
            "• Supercharger 운영·확장",
            "• 보급형 EV 플랫폼 개발",
        ],
        "label_color": WHITE,
        "bg": DARK,
    },
    "KR": {
        "label": "KR  |  핵심 자원",
        "lines": [
            "• 배터리 기술·수직계열화",
            "• FSD AI + 주행 데이터",
            "• 기가팩토리 (4개국)",
            "• Supercharger 네트워크",
            "• Tesla 브랜드·미션",
        ],
        "label_color": WHITE,
        "bg": DARK,
    },
    "VP": {
        "label": "VP  |  가치 제안",
        "lines": [
            "• 업계 최고 주행거리·성능",
            "• OTA 업데이트 (차가 진화)",
            "• FSD 완전자율주행",
            "• TCO 절감 (연료·유지비↓)",
            "• 지속가능성 미션 동참",
            "• 보급형: 경제적 합리성",
        ],
        "label_color": WHITE,
        "bg": RED,
    },
    "CR": {
        "label": "CR  |  고객 관계",
        "lines": [
            "• 직접 판매 (딜러 없음)",
            "• OTA로 지속적 관계 유지",
            "• Tesla 오너 커뮤니티",
            "• FSD 구독 반복 관계",
            "• 커뮤니티 교육 프로그램",
        ],
        "label_color": WHITE,
        "bg": DGRAY,
    },
    "CH": {
        "label": "CH  |  채널",
        "lines": [
            "• Tesla 직영 온라인 스토어",
            "• 직영 쇼룸 (도심)",
            "• Supercharger 네트워크",
            "• Tesla 앱 (통합 관리)",
            "• 모바일 서비스",
        ],
        "label_color": WHITE,
        "bg": DGRAY,
    },
    "CS": {
        "label": "CS  |  고객 세그먼트",
        "lines": [
            "[현재]",
            "• 고소득 ($125K+) 테크 얼리어답터",
            "[확장]",
            "• 중산층 가정",
            "• 저소득·유색인종 커뮤니티",
            "• 농촌·교외 거주자",
            "• 세입자 (미국 인구 36%)",
        ],
        "label_color": WHITE,
        "bg": BLUE,
    },
    "C$": {
        "label": "C$  |  비용 구조",
        "lines": [
            "• 배터리·원자재 (최대 비중)",
            "• 기가팩토리 건설·운영",
            "• R&D — FSD·배터리·신차",
            "• Supercharger 인프라",
            "• 보급형 플랫폼 개발비",
        ],
        "label_color": WHITE,
        "bg": DARK,
    },
    "R$": {
        "label": "R$  |  수익원",
        "lines": [
            "• 차량 판매 (주수익, 마진 압박 중)",
            "• FSD 구독 (고마진·반복 수익) ★",
            "• 에너지 사업 (Powerwall) ~$5B",
            "• Supercharger 충전 수익",
            "• Tesla 보험·금융 서비스",
            "• V2G 에너지 수익 (신규)",
        ],
        "label_color": WHITE,
        "bg": GREEN,
    },
}

for key, cfg in bmc_content.items():
    if key not in bmc_map:
        print(f"WARNING: shape '{key}' not found")
        continue
    shape = bmc_map[key]
    # Set background
    shape.fill.solid()
    shape.fill.fore_color.rgb = cfg["bg"]
    shape.line.fill.background()

    # Clear and rewrite text
    tf = shape.text_frame
    tf.word_wrap = True

    from pptx.oxml.ns import qn
    txBody = tf._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)

    all_items = [(cfg["label"], 9, True, cfg["label_color"], PP_ALIGN.LEFT, 4, 4)]
    for line in cfg["lines"]:
        sz = 7.5
        bold = line.startswith("[")
        color = RGBColor(0xFF,0xFF,0xAA) if line.startswith("[") else RGBColor(0xEE,0xEE,0xEE)
        all_items.append((line, sz, bold, color, PP_ALIGN.LEFT, 1, 0))

    first = True
    for (text, size, bold, color, align, sp_b, sp_a) in all_items:
        if first:
            p = tf.paragraphs[0] if tf.paragraphs else tf.add_paragraph()
            first = False
        else:
            p = tf.add_paragraph()
        p.text = text
        p.alignment = align
        p.space_before = Pt(sp_b)
        p.space_after  = Pt(sp_a)
        r = p.runs[0] if p.runs else p.add_run()
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color


# ── Save ───────────────────────────────────────────────
out = '/home/user/Business-Modeling/Tesla_Filled.pptx'
prs.save(out)
print(f"Saved: {out}")
